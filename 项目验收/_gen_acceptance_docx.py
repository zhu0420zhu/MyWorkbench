"""Generate Lingmo Workbench delivery & acceptance Word doc."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

OUT = Path(r"D:\YM\MyWorkbench\项目验收\灵墨工作台-项目交付与验收说明.docx")


def set_run_font(run, east="微软雅黑", latin="Calibri", size=11, bold=False, color=None):
    run.font.name = latin
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), east)
    rFonts.set(qn("w:ascii"), latin)
    rFonts.set(qn("w:hAnsi"), latin)


def add_para(doc, text, style=None, size=11, bold=False, color=None, space_after=6):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, size=10, bold=True)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = t.rows[r_i + 1].cells[c_i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            set_run_font(run, size=10)
    doc.add_paragraph()
    return t


def main():
    doc = Document()
    # A4
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(3.0)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("灵墨工作台")
    set_run_font(r, size=28, bold=True, color=(31, 58, 95))
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("项目交付说明与验收文档")
    set_run_font(r, size=18, bold=True, color=(31, 58, 95))
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run("版本 0.1.0 ｜ Windows x64 ｜ 生成日期 2026-09-19")
    set_run_font(r, size=11, color=(89, 89, 89))

    add_para(doc, "", space_after=8)

    # 0 Overview
    add_para(doc, "一、项目概览", style="Heading 1")
    add_para(
        doc,
        "灵墨工作台是一体化本地生产力工作台，集成快捷工具、知识库、写作（含长篇与工作流）、文档处理、"
        "文本批处理/OCR、数据整理与分析、白板、便利贴、情报热点、AI 对话等能力。数据以纯文本文件形式"
        "存储于本机，强调数据自主与离线可用；AI 采用 OpenAI 兼容接口，支持多套模型配置。",
    )
    add_table(
        doc,
        ["项目", "内容"],
        [
            ["产品名称", "灵墨工作台（Lingmo Workbench）"],
            ["版本号", "0.1.0"],
            ["运行平台", "Windows 10/11 x64"],
            ["交付形态", "NSIS 安装包 + 便携版"],
            ["技术栈", "Electron 43 + React 19 + TypeScript + Ant Design 5"],
            ["应用标识", "com.lingmo.workbench"],
            ["默认数据目录", "%APPDATA%\\lingmo-workbench\\（可自定义）"],
        ],
    )

    # 1 Features
    add_para(doc, "二、功能清单", style="Heading 1")
    add_para(doc, "2.1 侧栏模块", style="Heading 2")
    add_table(
        doc,
        ["模块", "主要功能", "验收重点"],
        [
            ["自动工作台", "快捷卡片、自动化命令、AI 入口、案例教程", "卡片与命令可运行"],
            ["本地知识库", "目录建索引、问答、数据资产、检索", "问答能引用资料"],
            ["文件工具", "磁盘/目录占用、重复文件、目录监控", "路径选择与结果"],
            ["文档处理", "MD/PDF/DOCX 预览编辑、Word 模式、打印、AI 右键", "DOCX 图文、保存"],
            ["写作", "短篇/长篇、工作流、角色/章节、AI 助手", "新建作品自动挂工作流"],
            ["文案", "文案场景对话，支持选中右键 AI", "右键菜单可用"],
            ["笔记", "Markdown、标签、快照历史、导出", "编辑保存与历史"],
            ["便利贴", "悬浮置顶、日历", "桌面小窗同步"],
            ["文本工具", "行处理/转换/清理/排版/编码/增强/替换/翻译/OCR/Diff", "分类工具与 OCR"],
            ["数据整理", "数据表、看板、可视化、多格式导入导出", "CSV 导入"],
            ["白板", "tldraw 无限画布", "绘制与自动保存"],
            ["情报中心", "热点聚合、推送、多平台分发、内嵌浏览", "拉取与打开事件"],
            ["数据分析", "图表、多维筛选、规则化结论", "筛选后图表变化"],
            ["设置", "存储、AI 多配置、个性化壁纸、MCP、数据看板", "配置与看板"],
        ],
    )

    add_para(doc, "2.2 重点能力说明", style="Heading 2")
    add_para(doc, "（1）写作与工作流", bold=True)
    add_para(
        doc,
        "新建长篇作品时按题材自动挂载写作工作流（默认「网文快写流」，可切换雪花创作法、三幕结构、"
        "英雄之旅、去 AI 味润色等）。工作流分步引导（开书三件套、金手指、分卷大纲、黄金三章、正文循环等），"
        "可一键调用 AI 生成，并勾选完成进度；配套写作 Skill 与灵感面板。",
    )
    add_para(doc, "（2）文档处理", bold=True)
    add_para(
        doc,
        "支持打开 Markdown、文本、PDF、DOCX、HTML、JSON、CSV、代码与图片。DOCX 经 Pandoc 预览并内嵌图片，"
        "提供 Word 视图（字体/字号/颜色/列表/对齐/行距/表格等）、查找替换、撤销重做、保存与另存、打印；"
        "DOCX 保存回写前生成 .bak 备份。",
    )
    add_para(doc, "（3）AI 能力", bold=True)
    add_para(
        doc,
        "OpenAI 兼容接口；设置中可维护多套模型配置（智谱、小米 MiMo、DeepSeek、通义等预设），一键启用切换。"
        "API Key 使用系统 safeStorage 加密存储。写作、文案、笔记、文档处理等场景支持选中文本后右键 AI"
        "（改写/扩写/缩写/润色/场景增强/去 AI 味/改基调/追问/总结/翻译）。",
    )
    add_para(doc, "（4）个性化与界面", bold=True)
    add_para(
        doc,
        "深浅双主题；氛境壁纸可选预设或自定义，并可对标题栏、侧边栏、内容区分别开关与配置；"
        "支持上传分区背景图；设置页分区卡片默认可折叠展开。无边框窗口与独立窗口打开页面。",
    )
    add_para(doc, "（5）数据与统计", bold=True)
    add_para(
        doc,
        "数据文件化存储于 userData，可备份/导入；设置内「数据看板」展示 AI 模型调用次数、估算 token、"
        "各模块内容量、页面访问次数等。",
    )

    # 2 Architecture
    add_para(doc, "三、项目架构摘要", style="Heading 1")
    add_table(
        doc,
        ["层次", "说明"],
        [
            ["界面层", "React + Ant Design；页面路由与侧栏由 navItems 维护"],
            ["桥接层", "Electron preload contextBridge（window.api）"],
            ["主进程", "窗口管理、IPC（domain:action）、业务存储与工具"],
            ["AI", "chat 流式 / complete；多配置；用量统计"],
            ["存储", "JSON/Markdown 文件 + 可选 SQLite 数据表 + KB 表库"],
            ["文档", "Pandoc 转换；OCR（WASM 离线 PP-OCRv4 + 表格模型）"],
        ],
    )
    add_para(
        doc,
        "源码主目录：src/main、src/preload、src/shared、src/renderer；项目验收文档位于仓库「项目验收」目录。",
    )

    # 3 Acceptance
    add_para(doc, "四、验收标准（摘要）", style="Heading 1")
    add_para(
        doc,
        "验收以「可安装、核心路径可用、数据可落盘、无阻断级崩溃」为准。建议按下列 P0 项逐条打勾，"
        "详细步骤见同目录《03-验收用例.md》。",
    )
    add_table(
        doc,
        ["编号", "验收项", "通过标准", "结果"],
        [
            ["P0-1", "安装启动", "安装包或便携版可启动，侧栏完整", "☐"],
            ["P0-2", "AI 配置", "可保存并测试连接；可多套切换；Key 非明文", "☐"],
            ["P0-3", "写作工作流", "新建作品自动挂工作流，可运行步骤", "☐"],
            ["P0-4", "文档处理", "打开 DOCX 图文正常，可编辑保存，右键 AI", "☐"],
            ["P0-5", "文本工具", "分类工具可用，OCR/Diff 可进入", "☐"],
            ["P0-6", "个性化", "壁纸/主题/设置折叠卡片生效", "☐"],
            ["P0-7", "数据看板", "调用后统计数字增长", "☐"],
        ],
    )

    add_para(doc, "4.1 建议验收流程（约 15 分钟）", style="Heading 2")
    steps = [
        "安装或运行最新安装包/便携版；",
        "设置 → AI 模型：添加配置并启用，点击测试连接；",
        "写作 → 新建作品（如题材「玄幻」）→ 确认右侧工作流面板；",
        "文档处理 → 打开测试 DOCX/MD → 编辑保存；",
        "文本工具 → 去重/全半角/词频试用；",
        "设置 → 氛境壁纸与数据看板检查；",
        "任意页面选中文本 → 右键验证 AI 菜单。",
    ]
    for i, s in enumerate(steps, 1):
        add_para(doc, f"{i}. {s}", space_after=3)

    # 4 Deliverables
    add_para(doc, "五、交付物", style="Heading 1")
    add_table(
        doc,
        ["交付物", "路径/说明"],
        [
            ["安装包", "dist\\灵墨工作台 Setup 0.1.0.exe（NSIS）"],
            ["便携版", "dist\\灵墨工作台 0.1.0.exe"],
            ["验收文档", "项目验收\\*.md 与本 Word/PDF"],
            ["源码与说明", "仓库根目录 README.md 等"],
        ],
    )
    add_para(
        doc,
        "打包说明：使用 electron-builder；建议打包前关闭正在运行的应用。"
        "Electron 相关下载建议配置 npmmirror 镜像（见 README）。",
    )

    # 5 Known issues
    add_para(doc, "六、已知问题与后续建议", style="Heading 1")
    add_table(
        doc,
        ["类型", "说明"],
        [
            ["工程", "Windows 下 electron-builder 偶发解压锁文件（EPERM），可清理 dist\\win-unpacked 后重试"],
            ["依赖", "Excel 解析库 xlsx@0.18.5 版本偏旧，建议后续替换"],
            ["功能", "自动更新未接入；知识库检索为关键词倒排（非向量）"],
            ["文档", "DOCX 复杂样式往返保存可能有差异，重要稿件建议另存 Word 后微调"],
            ["路线", "建议：自动更新、exceljs、向量检索、繁简转换、看板日报曲线等"],
        ],
    )

    # 7 Sign-off
    add_para(doc, "七、验收结论", style="Heading 1")
    add_table(
        doc,
        ["项目", "填写"],
        [
            ["P0 通过数", "______ / 7"],
            ["验收结论", "□ 通过　□ 有条件通过　□ 不通过"],
            ["问题清单", "（可另附截图与复现步骤）"],
            ["验收人", "________________"],
            ["日期", "____年__月__日"],
        ],
    )
    add_para(
        doc,
        "备注：本说明为交付与验收用汇总文档；更细的功能对照与用例步骤以「项目验收」目录下 Markdown 为准。",
        size=10,
        color=(89, 89, 89),
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print("WROTE", OUT)


if __name__ == "__main__":
    main()
